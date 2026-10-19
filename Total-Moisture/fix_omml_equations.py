import docx

doc = docx.Document('_ Draft.docx')

# Fix P49 (Free moisture)
p49 = doc.paragraphs[49]
num = p49._p.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}num')
if num is not None:
    subs = num.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}sub')
    if len(subs) >= 2:
        t0 = subs[0].find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}t')
        t1 = subs[1].find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}t')
        if t0 is not None and t1 is not None:
            print(f"P49 numerator before: {t0.text} - {t1.text}")
            t0.text = "1"
            t1.text = "2"
            print(f"P49 numerator after: {t0.text} - {t1.text}")

# Fix P54 (Residual moisture)
p54 = doc.paragraphs[54]
num54 = p54._p.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}num')
if num54 is not None:
    subs54 = num54.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}sub')
    if len(subs54) >= 2:
        t0 = subs54[0].find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}t')
        t1 = subs54[1].find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}t')
        if t0 is not None and t1 is not None:
            print(f"P54 numerator before: {t0.text} - {t1.text}")
            t0.text = "3"
            t1.text = "4"
            print(f"P54 numerator after: {t0.text} - {t1.text}")

doc.save('_ Draft.docx')
doc.save('AS015IW DETERMINATION OF TOTAL MOISTURE Rev 01.docx')
print("Successfully fixed OMML equations and saved documents!")
