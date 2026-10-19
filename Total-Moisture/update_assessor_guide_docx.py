import sys
import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def update_assessor_guide():
    doc = docx.Document('AZCA-015 Assessor Guide and Marking Memo.docx')
    
    # Update Paragraphs
    doc.paragraphs[0].text = "MODEL ANSWERS AND MARKING RUBRIC (TOTAL: 25 POINTS)"
    doc.paragraphs[1].text = "ASSESSMENT DECISION RULES:"
    doc.paragraphs[2].text = "• Qualifying Score: Candidate must achieve ≥ 80% (minimum 20 / 25 points) AND 100% on both Critical Questions (*5 and *6) → COMPETENT."
    doc.paragraphs[3].text = "• Zero-Tolerance Failure: If Question *5 OR Question *6 is incorrect (even if total score ≥ 20) → NOT YET COMPETENT (Immediate assessment stoppage)."
    doc.paragraphs[4].text = "• Insufficient Score: If Candidate Score < 20 points → NOT YET COMPETENT."
    
    # Update Table 0
    tbl = doc.tables[0]
    
    # Column Headers: Q# | Question Focus | Model Answer & Mandatory Technical Criteria | Points
    rubric_data = [
        ("1", 
         "Method Selection & Principle\n(AS015IW Cl. 1, ISO 589 Cl. 4 & 7.5)", 
         "• Two methods: Single-Stage and Two-Stage thermal drying (1 pt).\n"
         "• Criteria for Two-Stage: Required when coal is visibly wet / excessively moist, where free water separates in the container (air-dried with container per ISO 13909-4 / ISO 589 Cl. 7.5), or when particle size/mass makes single-stage crushing prone to moisture loss or equipment choking (1 pt).", 
         "2 pts"),
         
        ("2", 
         "Sample Mass & Top-Size\n(AS015IW Cl. 2.1, ISO 13909-4 Tab. 1, AS001IW Cl. 4.1)", 
         "• Minimum bulk sample mass for 50 mm top-size after division: 35 kg per ISO 13909-4:2025 Table 1 (1 pt).\n"
         "• Maximum nominal top-size accepted by laboratory: Shall not be > 50 mm per AS001IW Clause 4.1 (1 pt).\n"
         "(Note: 35 kg is the minimum divided bulk sample mass, not the individual drying pan charge).", 
         "2 pts"),
         
        ("3", 
         "Cone & Quartering\n(AS015IW Cl. 2.1, ISO 13909-4)", 
         "• Procedure: Form coal into cone, flatten top uniformly, divide into 4 equal quarters, select one quarter (1 pt).\n"
         "• Action if selected quarter < 35 kg: Analyst must combine two opposite quarters to ensure statistical representativeness (1 pt).", 
         "2 pts"),
         
        ("4", 
         "Stage 1 Air-Drying Conditions & Rationale\n(AS015IW Cl. 2.2, ISO 589 Cl. 6.1.1 & 8.1.1)", 
         "• Conditions: Air dry at ambient temperature or in fan-forced oven not exceeding 40 °C (layer loading ≤ 1 g/cm²), and re-equilibrate to ambient before weighing if heated (1 pt).\n"
         "• Chemical Rationale: Temperatures > 40 °C promote rapid coal oxidation (especially in lower-rank coals), causing chemical oxygen chemisorption, mass gain, and severe negative analytical bias (1 pt).", 
         "2 pts"),
         
        ("*5", 
         "CRITICAL: Procedural Flow & Crushing\n(AS015IW Cl. 2.2, 3.1, 3.2 — ZERO TOLERANCE)", 
         "Mandatory 5 sequential steps (1 pt each — ALL must be correct in order):\n"
         "1. Spread 35 kg sample thinly on clean drying pans until 'visibly dry' (≤ 40 °C; loading ≤ 1 g/cm²).\n"
         "2. Stage 1 stability check on LIMS: air drying complete when mass loss < 0.1 g (ISO constant-mass: change ≤ 0.2% absolute).\n"
         "3. Primary crush through 25 mm jaw crusher and reduce mass using clean Riffle Splitter.\n"
         "4. Secondary crush: crush to -10 mm for Single-Stage, or -3 mm for Two-Stage (AzureSky validated macro-pan method).\n"
         "5. Final riffle division to 1 kg per side, tare pans on balance, transfer into duplicate Pan A and Pan B, and capture initial mass m3.", 
         "*5 pts\n(Critical Gate 1)"),
         
        ("*6", 
         "CRITICAL: Oven Temperature & Stability\n(AS015IW Cl. 4.1, ISO 589 Cl. 4 — ZERO TOLERANCE)", 
         "• Operating Temperature: 105 °C to 110 °C (nominal setpoint ≤ 107 °C) (1 pt).\n"
         "• Drying Sequence & Stability: Initial dry for 3 hours; intermediate weighing on LIMS, return to oven for 30 minutes. Stable when mass loss between consecutive 30-min intervals is ≤ 0.1 g (1 pt).\n"
         "(Note: ISO constant mass is ≤ 0.2% absolute over ≥ 25% duration; ≤ 0.1 g represents AzureSky operational LIMS threshold).", 
         "*2 pts\n(Critical Gate 2)"),
         
        ("7", 
         "Oven Failure Threshold & Mandatory Action\n(AS015IW Cl. 4.1.5)", 
         "• Failure Limit: If sample does not stabilize within 5.0 hours of total oven drying time (0.5 pt).\n"
         "• Mandatory Action: Test marked as 'Failed' in LIMS, test portion invalidated, and determination must be completely repeated using a backup split from -3 mm retained sample (0.5 pt).", 
         "1 pt"),
         
        ("8", 
         "Single-Stage vs Two-Stage Oven Drying\n(AS015IW Cl. 4.2, ISO 589 Method B2)", 
         "• Continuous Drying: In Single-Stage (-10 mm), duplicate test portions (≥ 600 g, routine ~1000 g) remain in oven for the full 5-hour continuous duration without the 3-hour lapse check, dried at 105–110 °C (0.5 pt).\n"
         "• Hot Weighing: Trays removed and weighed hot within 5 minutes per ISO 589 Clause 8.2.1 to prevent moisture re-adsorption from ambient air during cooling (0.5 pt).", 
         "1 pt"),
         
        ("9", 
         "Equipment Metrology & Verification Schedules\n(AS023TF, AS001IW Cl. 4.1)", 
         "• Moisture Oven Verification (AS023TF): Calibrated thermocouple reading (average of 3 readings + correction factor), evaluated against set point within tolerance 105 °C to 110 °C (setpoint ≤ 107 °C) daily before use (1 pt).\n"
         "• Crusher Sieve Efficiencies (AS001IW Cl. 4.1): 25 mm & 10 mm crushers verified every Monday at start of shift; 3 mm & 4.75 mm crushers verified on day of use before use (1 pt).", 
         "2 pts"),
         
        ("10", 
         "Surface Moisture Formula & Calculation\n(AS015IW Cl. 5.1, ISO 589 Eq. 1)", 
         "• Formula & Variables: Mar = [(m1 − m2) / m1] × 100, where m1 is initial as-received bulk mass (kg) and m2 is air-dried bulk mass (kg) (1 pt).\n"
         "• Worked Calculation: Given m1 = 35.00 kg, m2 = 32.55 kg:\n"
         "  Mar = [(35.00 − 32.55) / 35.00] × 100 = (2.45 / 35.00) × 100 = 7.00% (1 pt).", 
         "2 pts"),
         
        ("11", 
         "Repeatability Compliance & LIMS Decision Rule\n(AS015IW Cl. 6, ISO 589 Tab. 1, AS010IW)", 
         "• Repeatability Evaluation: |4.58% − 4.20%| = 0.38%. Since 0.38% > 0.30% (ISO 589 second-stage repeatability limit), the duplicate pair DOES NOT comply (1 pt).\n"
         "• LIMS Alert & Corrective Action: Under AzureSky decision rules, LIMS triggers an automated 'REPEAT ALERT'; results are locked, invalidated, and a repeat test from retained sample must be performed (1 pt).", 
         "2 pts"),
         
        ("12", 
         "LIMS vs ISO 589 Calculation Alignment\n(AS015IW Cl. 5.3 & 6)", 
         "• Calculation Flow Difference: LIMS averages duplicate residual moistures first (Mres,avg) and inserts directly into TM formula: TM = Mar + [(100 − Mar) × Mres,avg]/100. ISO 589 calculates TM of Pan A and Pan B individually and averages them at the end (1 pt).\n"
         "• Equivalence & QMS Approval: Because the formula is mathematically linear, both methods yield identical numerical results, and both are formally approved under AzureSky QMS (1 pt).", 
         "2 pts")
    ]
    
    # Populate rows 1 to 12
    for r_idx, (q_num, focus, answer, pts) in enumerate(rubric_data, start=1):
        row = tbl.rows[r_idx]
        row.cells[0].text = q_num
        row.cells[1].text = focus
        row.cells[2].text = answer
        row.cells[3].text = pts
        
    doc.save('AZCA-015 Assessor Guide and Marking Memo.docx')
    print("Successfully updated AZCA-015 Assessor Guide and Marking Memo.docx to 25-point responsible scheme!")

if __name__ == '__main__':
    update_assessor_guide()
