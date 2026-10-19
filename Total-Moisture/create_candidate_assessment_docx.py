import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_candidate_assessment():
    doc = docx.Document()
    
    # Page setup - Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Title Block Table
    t_header = doc.add_table(rows=2, cols=2)
    t_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_header.autofit = False
    
    cell_00 = t_header.rows[0].cells[0]
    p_logo = cell_00.paragraphs[0]
    r_logo = p_logo.add_run("AzurESky Coal Laboratory")
    r_logo.font.size = Pt(14)
    r_logo.font.bold = True
    r_logo.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_sub = cell_00.add_paragraph("Technical Competency Evaluation System\nStandard: ISO/IEC 17025:2017 Cl. 6.2")
    p_sub.runs[0].font.size = Pt(8.5)
    p_sub.runs[0].font.italic = True
    
    cell_01 = t_header.rows[0].cells[1]
    cell_01.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_docid = cell_01.paragraphs[0].add_run("Document: AZCA-015 Rev 01\nModule: Total Moisture (AS015IW Rev 01)\nGoverning Standard: ISO 589 / ISO 13909-4")
    r_docid.font.size = Pt(8.5)
    r_docid.font.bold = True
    
    cell_title = t_header.rows[1].cells[0]
    cell_title.merge(t_header.rows[1].cells[1])
    p_t = cell_title.paragraphs[0]
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("TECHNICAL COMPETENCY ASSESSMENT: DETERMINATION OF TOTAL MOISTURE\n(SINGLE-STAGE & TWO-STAGE THERMAL DRYING)")
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    # Add a border / shade to header table
    for row in t_header.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="0" w:color="003366"/></w:tcBorders>')
            tcPr.append(tcBorders)

    doc.add_paragraph()

    # Candidate Identification Block
    t_cand = doc.add_table(rows=3, cols=4)
    t_cand.alignment = WD_TABLE_ALIGNMENT.CENTER
    cand_fields = [
        ("Candidate Name:", "", "Staff Number:", ""),
        ("Assessment Date:", "", "Attempt Number:", "[  ] 1st Attempt   [  ] 2nd Re-Assessment"),
        ("Assessor Name:", "", "Assessor Title:", "Laboratory Supervisor / Tech Signatory")
    ]
    for r_idx, row in enumerate(cand_fields):
        for c_idx, val in enumerate(row):
            cell = t_cand.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx % 2 == 0:
                r.font.bold = True
                r.font.size = Pt(9.5)
            else:
                r.font.size = Pt(9.5)
                
    for row in t_cand.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/></w:tcBorders>')
            tcPr.append(tcBorders)

    doc.add_paragraph()

    # Instructions Box
    p_inst_h = doc.add_paragraph()
    r = p_inst_h.add_run("MANDATORY ASSESSMENT RULES & QUALIFICATION CRITERIA")
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x99, 0x00, 0x00)

    p_inst = doc.add_paragraph()
    p_inst.add_run(
        "1. Purpose & Outcome: This assessment verifies technical competence in conducting Single-Stage and Two-Stage Total Moisture determinations in coal per procedure AS015IW Rev 01, ISO 589:2008, and ISO 13909-4:2025.\n"
        "2. Total Marks & Pass Mark: The assessment consists of 25 points. To be declared COMPETENT, the candidate must achieve a minimum of 80% (20 / 25 points).\n"
        "3. CRITICAL PROCESS GATES (Zero Tolerance): Questions *5 and *6 represent critical process control points. Candidates must achieve 100% on both Question *5 and Question *6. If either question is incorrect, the assessment is STOPPED IMMEDIATELY and the outcome is declared NOT YET COMPETENT.\n"
        "4. Permitted Materials: Candidate may reference controlled test procedure AS015IW Rev 01 and raw data sheets AS015TF-01/02 during the assessment."
    )
    p_inst.runs[0].font.size = Pt(9.5)
    p_inst.runs[0].font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Questions List
    questions = [
        ("Question 1 [2 Marks]",
         "Name the two thermal drying methods described in AS015IW for total moisture determination. Under what physical sample conditions is a laboratory required to perform a Two-Stage determination rather than a Single-Stage determination?\n\n\n\n"),
         
        ("Question 2 [2 Marks]",
         "In accordance with ISO 13909-4:2025 Table 1 and AS001IW Clause 4.1:\n"
         "a) What is the minimum mass of the divided bulk sample required for a 50 mm nominal top-size coal sample?\n"
         "b) What is the maximum nominal top-size accepted by the AzureSky coal laboratory?\n\n\n"),
         
        ("Question 3 [2 Marks]",
         "Describe the Cone and Quartering primary division procedure. If the mass of the chosen single quarter is less than the mandatory minimum mass, what specific action must the analyst take to ensure representativeness?\n\n\n\n"),
         
        ("Question 4 [2 Marks]",
         "At what temperature may the sample be air-dried during Stage 1? Explain the scientific rationale: why must the drying temperature strictly NOT exceed 40 °C, and what must be done if a heated oven is used before re-weighing?\n\n\n\n"),
         
        ("Question *5 [5 Marks — CRITICAL GATE 1: ZERO TOLERANCE]",
         "List the 5 key procedural steps in exact sequential order from Stage 1 air drying through to the preparation and weighing of duplicate test pans:\n"
         "Step 1: _____________________________________________________________________________________\n"
         "Step 2: _____________________________________________________________________________________\n"
         "Step 3: _____________________________________________________________________________________\n"
         "Step 4: _____________________________________________________________________________________\n"
         "Step 5: _____________________________________________________________________________________\n"),
         
        ("Question *6 [2 Marks — CRITICAL GATE 2: ZERO TOLERANCE]",
         "State the exact operating temperature range for the moisture oven. In a Two-Stage determination (-3 mm), what is the initial drying duration, and what is the stability criterion between 30-minute intervals to declare the sample stable?\n\n\n\n"),
         
        ("Question 7 [1 Mark]",
         "What is the maximum allowable drying time in the oven before the test is marked as 'Failed' in LIMS? What mandatory action must be taken if a sample fails to stabilize within this time?\n\n\n"),
         
        ("Question 8 [1 Mark]",
         "How does the oven drying procedure for Single-Stage moisture (-10 mm) differ from Two-Stage moisture (-3 mm) regarding the initial 3-hour check? State the weighing requirement upon removing the hot tray from the oven.\n\n\n"),
         
        ("Question 9 [2 Marks]",
         "According to AS023TF and AS001IW Clause 4.1:\n"
         "a) How is the moisture drying oven verified daily, and what is the operating tolerance limit?\n"
         "b) On what schedule must crusher sieve efficiencies be checked for 25 mm/10 mm coarse crushers and 3 mm/4.75 mm fine crushers?\n\n\n\n"),
         
        ("Question 10 [2 Marks]",
         "Write the formula for calculating Surface (Free) Moisture percentage (Mar) with defined variables.\n"
         "Calculate Mar for a bulk sample received at m1 = 35.00 kg and air-dried to m2 = 32.55 kg (show all working):\n\n\n\n"),
         
        ("Question 11 [2 Marks]",
         "During duplicate testing of residual moisture, Pan A yields Mres = 4.20% and Pan B yields Mres = 4.58%.\n"
         "a) Does this duplicate pair comply with the ISO 589 repeatability limit? (State the limit and show calculation).\n"
         "b) What automated alert is triggered in LIMS, and what mandatory action must the analyst take?\n\n\n\n"),
         
        ("Question 12 [2 Marks]",
         "Explain how AzureSky LIMS calculates Total Moisture compared to how ISO 589 calculates it for duplicate pans. Are both methods accepted under AzureSky QMS, and why?\n\n\n\n")
    ]

    for q_title, q_body in questions:
        p_q = doc.add_paragraph()
        r = p_q.add_run(q_title)
        r.font.bold = True
        r.font.size = Pt(10.5)
        if 'CRITICAL' in q_title:
            r.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
        else:
            r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            
        p_b = doc.add_paragraph(q_body)
        p_b.paragraph_format.line_spacing = 1.15
        p_b.paragraph_format.space_after = Pt(8)

    # Sign-Off Table
    doc.add_page_break()
    p_dec_h = doc.add_paragraph()
    r_dec = p_dec_h.add_run("ASSESSMENT OUTCOME, FEEDBACK & DECLARATION SIGN-OFF")
    r_dec.font.bold = True
    r_dec.font.size = Pt(11)
    r_dec.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    t_dec = doc.add_table(rows=7, cols=2)
    t_dec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_dec.autofit = False
    
    t_dec.rows[0].cells[0].text = "Total Candidate Score:"
    t_dec.rows[0].cells[1].text = "_______ / 25 Marks ( _____ % )"
    
    t_dec.rows[1].cells[0].text = "Critical Gates Status:"
    t_dec.rows[1].cells[1].text = "Question *5 (Procedural Flow): [  ] PASS   [  ] FAIL\nQuestion *6 (Temp & Timing):   [  ] PASS   [  ] FAIL"
    
    t_dec.rows[2].cells[0].text = "Competency Determination:"
    t_dec.rows[2].cells[1].text = "[   ] COMPETENT (Score ≥ 20/25 AND both Critical Gates Passed)\n[   ] NOT YET COMPETENT (Score < 20 OR any Critical Gate Failed)"
    
    t_dec.rows[3].cells[0].text = "Assessor Technical Comments & Recommendations:"
    t_dec.rows[3].cells[1].text = "\n\n\n"
    
    t_dec.rows[4].cells[0].text = "Candidate Feedback & Acknowledgement:"
    t_dec.rows[4].cells[1].text = "\n\n\n"
    
    t_dec.rows[5].cells[0].text = "Assessor Declaration:\nI hereby certify that the candidate has been assessed in accordance with AZCA-015 and ISO/IEC 17025:2017 Clause 6.2."
    t_dec.rows[5].cells[1].text = "Assessor Name: __________________________\nSignature:     __________________________\nDate:          __________________________"
    
    t_dec.rows[6].cells[0].text = "Candidate Declaration:\nI declare that I have received feedback on my assessment and understand the outcome."
    t_dec.rows[6].cells[1].text = "Candidate Name: __________________________\nSignature:      __________________________\nDate:           __________________________"

    for row in t_dec.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="003366"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="003366"/><w:left w:val="single" w:sz="4" w:space="0" w:color="003366"/><w:right w:val="single" w:sz="4" w:space="0" w:color="003366"/></w:tcBorders>')
            tcPr.append(tcBorders)
            
    doc.save('AZCA-015 Assessment for Total Moisture.docx')
    print("Successfully generated AZCA-015 Assessment for Total Moisture.docx candidate test paper!")

if __name__ == '__main__':
    create_candidate_assessment()
