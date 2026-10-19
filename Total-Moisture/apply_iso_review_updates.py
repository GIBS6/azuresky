import sys
import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

def apply_iso_review_updates():
    doc = docx.Document('_ Draft.docx')

    # =========================================================================
    # 1. Update Running Footer
    # =========================================================================
    for section in doc.sections:
        footer = section.footer
        for table in footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            if '00' in r.text:
                                r.text = r.text.replace('00', '01')
                            if '10/2025' in r.text:
                                r.text = r.text.replace('10/2025', '10/2026')
                            if '10/2028' in r.text:
                                r.text = r.text.replace('10/2028', '10/2029')

    # =========================================================================
    # 2. Iterate and update paragraphs
    # =========================================================================
    for p in doc.paragraphs:
        t = p.text.strip()

        # ---------------------------------------------------------------------
        # Clause 1: Scope, Principle & Method Selection
        # ---------------------------------------------------------------------
        if t.startswith('This method describes the determination of total moisture'):
            p.text = (
                "This procedure describes the determination of total moisture in hard coal using thermal drying in air. "
                "The procedure specifies an AzureSky In-House Validated Method based on the principles of ISO 589:2008 and "
                "ISO 13909-4:2025, operating under AzureSky's ISO/IEC 17025 management system.\n\n"
                "Moisture is determined gravimetrically by the loss in mass of the sample when heated in air at 105 ºC to 110 °C "
                "under controlled conditions until constant mass is attained. This method is applicable to hard coals not susceptible "
                "to significant thermal oxidation."
            )
        elif 'Reference Standards:' in t and 'Method Selection Criteria:' in t:
            p.text = (
                "Governing Standards & References: ISO 589:2008 (Hard coal — Determination of total moisture), "
                "ISO 13909-4:2025 (Coal — Preparation of test samples), ISO 18283, AS001IW, AS023TF.\n\n"
                "Method Selection & Scope Criteria:\n"
                "• Two-Stage Determination (Clause 2, 3 & 4.1): Applied to unreduced bulk samples, visibly wet coals, "
                "or excessively moist consignments where direct mechanical reduction would cause moisture loss by evaporation, "
                "choking of crushers, or screen blinding. Where free water separates from the coal in the sample container, the whole "
                "sample and its container shall be air-dried together in accordance with ISO 13909-4:2025 and ISO 589 Clause 7.5, "
                "and the mass loss accounted for.\n"
                "• Single-Stage Determination (Clause 4.2): Applied to free-flowing, drier coals received or prepared directly "
                "at nominal top-size -10 mm."
            )

        # ---------------------------------------------------------------------
        # Clause 2: Sample Preparation and Division
        # ---------------------------------------------------------------------
        elif t.startswith('1. Sample Acceptance:'):
            p.text = (
                "1. Sample Acceptance & Bulk Mass Requirement: The nominal top-size of incoming coal accepted by the laboratory "
                "shall not exceed 50 mm (AS001IW Cl. 4.1). In accordance with ISO 13909-4:2025 Table 1, the minimum required mass "
                "of the divided bulk sample for a 50 mm nominal top-size is 35 kg. Note: This 35 kg constitutes the minimum sample "
                "mass required after division to maintain sampling precision; it does not represent the mass placed on an individual test pan."
            )
        elif t.startswith('3. Division & Representativeness:'):
            p.text = (
                "3. Division & Representativeness: Select one quarter. If the mass of the chosen single quarter is less than the mandatory "
                "35 kg minimum sample mass, combine two opposite quarters to ensure statistical representativeness."
            )
        elif t.startswith('4. Initial Mass (m_1):'):
            p.text = (
                "4. Initial Bulk Mass (m_1): Immediately weigh the divided bulk sample on the calibrated platform balance and record the mass "
                "on LIMS and raw data sheet AS015TF-02. This is the initial \"As Received\" bulk mass (m_1, in kg)."
            )

        # ---------------------------------------------------------------------
        # Clause 2.2: Stage 1 Air Drying
        # ---------------------------------------------------------------------
        elif t.startswith('1. Spreading:'):
            p.text = (
                "1. Spreading: Spread the divided sample thinly and evenly over clean drying pans such that the coal layer loading "
                "does not exceed 1 g/cm² (ISO 589 Clause 6.3.1)."
            )
        elif t.startswith('2. Drying Conditions:'):
            p.text = (
                "2. Drying Conditions & Oxidation Safeguard: Allow the sample to air dry at ambient laboratory temperature or in a "
                "fan-forced oven not exceeding 40 °C (ISO 589 Clause 8.1.1) until visibly dry.\n"
                "• Temperature Restriction: The drying temperature must strictly NOT exceed 40 °C. Temperatures above 40 °C promote "
                "thermal oxidation (especially in lower-rank and bituminous coals), resulting in chemical oxygen chemisorption, weight gain, "
                "and severe negative analytical bias.\n"
                "• Re-equilibration: When a heated oven (≤ 40 °C) is used for air drying, the coal sample must re-equilibrate to ambient "
                "room temperature before re-weighing (ISO 589 Clause 8.1.1)."
            )
        elif t.startswith('3. Stability Check:'):
            p.text = (
                "3. Stage 1 Constant-Mass & Stability Criteria:\n"
                "• ISO 589 Normative Definition: Constancy in mass is reached when the change in mass does not exceed 0.2% absolute "
                "during a further period of not less than 25% of the initial drying period.\n"
                "• AzureSky Operational Stability Criterion: In daily routine laboratory operation on LIMS, air drying is considered complete "
                "when the mass loss between successive check weighings is less than 0.1 g (< 0.1 g)."
            )
        elif t.startswith('4. Final Air-Dried Mass (m_2):'):
            p.text = (
                "4. Final Air-Dried Bulk Mass (m_2): Record the final air-dried bulk mass (m_2, in kg) on LIMS and AS015TF-02."
            )

        # ---------------------------------------------------------------------
        # Clause 3: Sub-Sampling and Crushing
        # ---------------------------------------------------------------------
        elif t.startswith('3. Sub-Sampling and Crushing'):
            p.text = "3. Sub-Sampling and Crushing Sequence (ISO 13909-4 / AzureSky Method)"
        elif t.startswith('• Step 3 (Secondary Crush):'):
            p.text = (
                "• Step 3 (Secondary Comminution):\n"
                "    - For Single-Stage: Crush sample to nominal top-size -10 mm.\n"
                "    - For Two-Stage (AzureSky Validated Procedure): Crush sample to -3 mm using the calibrated jaw crusher.\n"
                "    * Technical Note on ISO 589 Method A2 Alignment: ISO 589:2008 Clause 8.1.2 specifies reduction to 2.8 mm nominal "
                "top-size with micro test portions of (10 ± 1) g in shallow dishes (layer loading ≤ 0.3 g/cm²). AzureSky's accredited macro "
                "procedure uses 1 kg test portions crushed to -3 mm in drying pans as an in-house validated adaptation to enhance "
                "sample representativeness on heterogeneous bulk coals."
            )

        # ---------------------------------------------------------------------
        # Clause 3.2: Final Preparation
        # ---------------------------------------------------------------------
        elif t.startswith('• Transfer & Initial Mass (m_3):'):
            p.text = (
                "• Transfer & Initial Sample Charge (m_3): Transfer approximately 1 kg into Pan A and 1 kg into Pan B. "
                "Immediately capture the net initial sample mass (m_3, in grams) on LIMS and AS015TF-02 for both duplicate pans."
            )

        # ---------------------------------------------------------------------
        # Clause 4: Oven Drying
        # ---------------------------------------------------------------------
        elif t.startswith('4. Oven Drying'):
            p.text = "4. Oven Drying (Residual and Single-Stage Moisture)"
        elif t.startswith('4.1. Two-Stage Moisture (-3 mm)'):
            p.text = "4.1. Two-Stage Residual Moisture (-3 mm AzureSky Macro Procedure)"
        elif t.startswith('1. Place duplicate pans in the oven pre-heated'):
            p.text = (
                "1. Oven Operating Conditions: Place duplicate pans in the oven controlled at 105 °C to 110 °C. The oven ventilation "
                "must maintain an atmosphere change rate of at least 5 times per hour per ISO 589 Clause 6.1.3."
            )
        elif t.startswith('4. Stability Criteria:'):
            p.text = (
                "4. Constant-Mass & Stability Criteria:\n"
                "• ISO 589 Normative Constant-Mass: Mass change not exceeding 0.2% absolute over a further drying period of not less than "
                "25% of the initial period.\n"
                "• AzureSky Operational Stability Criterion: On LIMS, the sample status reads \"stable\" when the mass loss between "
                "successive 30-minute drying intervals is ≤ 0.1 g. Once stable, record the final net dried mass (m_4, in grams) on LIMS."
            )
        elif t.startswith('5. Failure Limit:'):
            p.text = (
                "5. Maximum Oven Limit & Failure Action: If the sample does not stabilize within 5 hours of total drying, the test is marked "
                "as \"Failed\" in LIMS. The test is invalidated and must be completely repeated using a backup split."
            )

        # ---------------------------------------------------------------------
        # Clause 4.2: Single-Stage Moisture (-10 mm)
        # ---------------------------------------------------------------------
        elif t.startswith('1. Duplicate 1 kg portions crushed to -10 mm'):
            p.text = (
                "1. Sample Portions & Loading: Extract at least two duplicate test portions of minimum 600 g each (AzureSky routinely uses ~1000 g) "
                "from the -10 mm crushed sample (ISO 589 Clause 8.2.1). Spread evenly on clean trays (loading ≤ 1 g/cm²).\n"
                "2. Drying Conditions: Dry at 105 °C to 110 °C until constant mass is attained (ISO 589 Method B2). In AzureSky routine operations, "
                "a continuous baseline duration of 5.0 hours is applied, confirmed by constant mass.\n"
                "3. Hot Weighing: In strict compliance with ISO 589 Clause 8.2.1, remove trays from the oven and weigh while hot as soon as possible "
                "(within 5 minutes) to avoid moisture adsorption from ambient air during cooling."
            )
        elif t.startswith('4. Capture final stabilized mass on LIMS and AS015TF-01.'):
            p.text = (
                "4. Mass Capture & Records: Record empty tray tare (m_1), gross initial mass before drying (m_2), and gross dried mass (m_3) "
                "on LIMS and raw data sheet AS015TF-01."
            )

        # ---------------------------------------------------------------------
        # Clause 5: Calculations (Variable Separations)
        # ---------------------------------------------------------------------
        elif t.startswith('5.1. Surface (Free) Moisture (M_ad)'):
            p.text = "5.1. Surface (Free) Moisture (M_ar) — Two-Stage Stage 1"
        elif t.startswith('m_1 = Initial mass of sample as received.'):
            p.text = "m_1 = Initial mass of bulk sample as received (kg, per AS015TF-02)."
        elif t.startswith('m_2 = Mass of sample after air drying (kg).'):
            p.text = (
                "m_2 = Mass of bulk sample after Stage 1 air drying (kg, per AS015TF-02).\n\n"
                "Calculation Example (Assessment Q10):\n"
                "Given m_1 = 35.00 kg (as received) and m_2 = 32.55 kg (air dried):\n"
                "M_ar = [(35.00 - 32.55) / 35.00] × 100 = (2.45 / 35.00) × 100 = 7.00%"
            )
        elif t.startswith('5.2. Residual Moisture (M_res)'):
            p.text = "5.2. Residual Moisture (M_res) — Two-Stage Stage 2"
        elif t.startswith('m_3 = Mass of sample before oven drying (1 kg).'):
            p.text = "m_3 = Net sample mass before oven drying (~1000 g net charge, per AS015TF-02)."
        elif t.startswith('m_4 = Mass of sample after oven drying (stabilized).'):
            p.text = "m_4 = Final net stabilized sample mass after oven drying (g, per AS015TF-02)."

        # ---------------------------------------------------------------------
        # Clause 5.3: Total Moisture Equations
        # ---------------------------------------------------------------------
        elif t.startswith('For a two-stage determination, the moisture is calculated as follows:'):
            p.text = (
                "5.3. Total Moisture (TM) Equations\n\n"
                "A. Two-Stage Total Moisture (ISO 589 Equation 3):\n"
                "TM = M_ar + [(100 - M_ar) × M_res] / 100\n"
                "Where M_ar is surface moisture (%) and M_res is residual moisture (%).\n\n"
                "B. Single-Stage Total Moisture (ISO 589:2008 Clause 8.2.3 Equation 5):\n"
                "TM = [(m_2 - m_3) / (m_2 - m_1)] × 100\n\n"
                "Where (strictly standardized per AS015TF-01 and ISO 589 Cl. 8.2.3):\n"
                "• m_1 = Mass of empty, clean, dry drying tray tare (g)\n"
                "• m_2 = Gross mass of drying tray plus test sample before oven drying (g)\n"
                "• m_3 = Gross mass of drying tray plus test sample after oven drying (g)\n"
                "Note: In Single-Stage determination, moisture loss is (m_2 - m_3) and initial test sample net mass is (m_2 - m_1). "
                "These variables are distinct from the Two-Stage bulk variables m_1 and m_2."
            )

        # ---------------------------------------------------------------------
        # Clause 6: Precision Limits vs AzureSky Retest Decision Rules
        # ---------------------------------------------------------------------
        elif t.startswith('Repeatability Limit & LIMS Repeat Alert:'):
            p.text = (
                "6.1. ISO 589 Normative Precision Limits vs. AzureSky Decision Rules\n\n"
                "A. ISO 589:2008 Normative Precision (Table 1):\n"
                "• Second-Stage Moisture Repeatability (r): 0.30% (same operator, same apparatus, duplicate determinations).\n"
                "• Total Moisture Repeatability (r): 0.50% (same laboratory, duplicate determinations).\n"
                "• Total Moisture Reproducibility Critical Difference (R): 1.50% (different laboratories).\n\n"
                "B. AzureSky Quality Decision Rule (ISO/IEC 17025 Clause 7.8.6 / AS010IW):\n"
                "To ensure reporting validity, AzureSky applies an internal automated decision rule in LIMS: "
                "If the difference between duplicate Pan A and Pan B exceeds 0.30% for second-stage moisture (|M_res,A - M_res,B| > 0.30%) "
                "or 0.50% for total moisture, LIMS triggers a \"REPEAT ALERT\". The test result is invalidated, quarantined from release, "
                "and a complete repeat determination is initiated.\n\n"
                "Worked Example (Assessment Q11):\n"
                "Duplicate residual moisture results: Pan A M_res = 4.20%, Pan B M_res = 4.58%.\n"
                "Difference: |4.58% - 4.20%| = 0.38%. Since 0.38% > 0.30%, the duplicate pair exceeds the repeatability limit. "
                "LIMS triggers a REPEAT ALERT and the sample must be re-tested."
            )

        # ---------------------------------------------------------------------
        # Clause 7: Equipment Verification and Quality Control
        # ---------------------------------------------------------------------
        elif t.startswith('To ensure metrological traceability under ISO/IEC 17025 Clause 6.4:'):
            p.text = (
                "To ensure metrological traceability and measurement validity under ISO/IEC 17025:2017 Clause 6.4 and Clause 7.7, "
                "equipment verification must strictly adhere to the approved laboratory control forms:\n\n"
                "1. Moisture Drying Oven (Ref: AS023TF Rev 00):\n"
                "• Frequency: Daily before use.\n"
                "• Verification Method: Calibrated thermocouple reading (Average of 3 readings + Thermocouple Correction Factor) "
                "evaluated against the oven indicator set point and verified within the working tolerance limit of 105 °C to 110 °C.\n\n"
                "2. Crusher Performance & Sieve Efficiencies (Ref: AS001IW Rev 04 Clause 4.1):\n"
                "• Coarse Crushers (25 mm and 10 mm): Efficiency check conducted every Monday at the beginning of the shift before work commences.\n"
                "• Fine Crushers (3 mm and 4.75 mm): Efficiency check conducted on the day of use, immediately before using the crusher."
            )

    # =========================================================================
    # 3. Update Revision Table (Table 1)
    # =========================================================================
    rev_table = doc.tables[1]
    row_rev1 = rev_table.rows[2]
    row_rev1.cells[0].text = "01"
    row_rev1.cells[1].text = "10/2026"
    row_rev1.cells[2].text = (
        "Technical revision and standardization:\n"
        "1. Standardized variable notation: isolated Single-Stage variables (m1=tare, m2=gross initial, m3=gross dried per ISO 589 Cl. 8.2.3 and AS015TF-01) from Two-Stage variables.\n"
        "2. Formally documented AzureSky in-house validated macro-pan procedure (1 kg at -3 mm) alongside ISO 589:2008 normative criteria.\n"
        "3. Aligned bulk preparation with ISO 13909-4:2025 Table 1 (35 kg min mass at 50 mm top-size) and documented free water separation handling.\n"
        "4. Incorporated ISO 589 constant-mass definitions (≤ 0.2% over ≥ 25% duration) and ambient re-equilibration requirements for air drying (≤ 40 °C).\n"
        "5. Differentiated ISO 589 normative precision limits from AzureSky internal LIMS repeat decision rules (AS010IW).\n"
        "6. Aligned equipment controls with AS023TF and AS001IW."
    )

    # Save to AS015IW DETERMINATION OF TOTAL MOISTURE Rev 01.docx and _ Draft_UPDATED.docx
    doc.save('AS015IW DETERMINATION OF TOTAL MOISTURE Rev 01.docx')
    doc.save('_ Draft_UPDATED.docx')
    print("Saved AS015IW DETERMINATION OF TOTAL MOISTURE Rev 01.docx and _ Draft_UPDATED.docx successfully!")
    
    try:
        doc.save('_ Draft.docx')
        print("Successfully updated _ Draft.docx directly!")
    except PermissionError:
        print("Note: _ Draft.docx is currently open in Microsoft Word by the user. Saved changes to _ Draft_UPDATED.docx and AS015IW DETERMINATION OF TOTAL MOISTURE Rev 01.docx so no work is overwritten or blocked!")

if __name__ == '__main__':
    apply_iso_review_updates()
